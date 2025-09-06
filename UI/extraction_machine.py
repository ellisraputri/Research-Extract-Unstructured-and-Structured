import numpy as np
import json, time, os, re
import pandas as pd
from docx import Document
from typing import Sequence
from dotenv import load_dotenv
from docx import Document
from langchain.output_parsers import PydanticOutputParser
from langchain.prompts import PromptTemplate
from pydantic import BaseModel
from tkinter import filedialog, messagebox
from google import genai
from google.genai import types
import ast

load_dotenv(".env")
API_KEY = os.getenv("GEMINI_API_KEY")
print(API_KEY)

client = genai.Client(
    api_key=API_KEY
)

def parse_age(s):
    patt = re.compile(r'(?:(\d+)\s*thn)?\s*(?:(\d+)\s*bln)?\s*(?:(\d+)\s*hr)?', re.I)
    t, b, h = patt.fullmatch(s.strip()).groups()
    return int(t or 0), int(b or 0), int(h or 0)

def generate_available_text(text):
    with open("columns.json", "r") as f:
        config = json.load(f)

    column_fields = config["column_fields"]

    attributes = {field: str for field in column_fields} 
    Columns = type("Columns", (BaseModel,), {"__annotations__": attributes})

    class Information(BaseModel):
        information: Sequence[Columns]

    parser = PydanticOutputParser(pydantic_object=Information)

    prompt = PromptTemplate(
        template="Ekstrak informasi berikut. Anda harus selalu mengembalikan JSON yang valid yang dipagari oleh blok kode markdown. Jangan kembalikan teks tambahan apa pun.:\n{format_instructions}\n{text}\n",
        input_variables=["text"],
        partial_variables={"format_instructions": parser.get_format_instructions()},
    )

    _input = prompt.format(text=text)

    model = "gemini-2.0-flash"
    contents = [
        types.Content(
            role="user",
            parts=[
                types.Part.from_text(text=_input),
            ],
        ),
    ]
    generate_content_config = types.GenerateContentConfig(
        response_mime_type="text/plain",
    )

    response_text=""
    for chunk in client.models.generate_content_stream(
        model=model,
        contents=contents,
        config=generate_content_config,
    ):
        response_text += chunk.text  

    print("Full response:", response_text)

    # Parse the response
    try:
        xdata = json.loads(response_text.strip("```json").strip("```"))  # Clean and parse JSON
        xdf = pd.DataFrame(xdata["information"])
    except (json.JSONDecodeError, KeyError) as e:
        print(f"Error parsing response: {e}")
        return pd.DataFrame()

    print(xdf.columns)
    return xdf

def parse_O_details(df_input):
    df = df_input.copy()
    if "detailPemeriksaan_extracted" not in df.columns:
        df["detailPemeriksaan_extracted"] = None

    row_idx = 0
    row = df.iloc[row_idx]

    raw_text = row.get("detailPemeriksaan", None)
    if pd.isna(raw_text) or not isinstance(raw_text, str) or raw_text.strip() == "":
        return df
    
    print("raw text", raw_text)

    prompt_text = f"""
        Ekstrak 'nadi', 'suhu', 'pernapasan', 'SPO2', 'tinggiBadan', 'beratBadan', 'tekananDarah' dari teks di bawah ini.
        Berikan hasilnya dalam format JSON tanpa tambahan teks atau penjelasan.

        Teks:
        \"{raw_text.strip()}\" 
    """

    try:
        model = "gemini-2.0-flash"
        contents = [
            types.Content(
                role="user",
                parts=[types.Part.from_text(text=prompt_text)],
            ),
        ]
        generate_content_config = types.GenerateContentConfig(
            response_mime_type="text/plain",
        )

        response_text = ""
        for chunk in client.models.generate_content_stream(
            model=model,
            contents=contents,
            config=generate_content_config,
        ):
            response_text += chunk.text

        print("resp", response_text)

        clean_text = re.sub(r"```json|```", "", response_text).strip()
        xdata = json.loads(clean_text)

        print("xdwfs", xdata)

        df.at[row_idx, "detailPemeriksaan_extracted"] = xdata

    except Exception as e:
        print("row error in O", e)
        df.at[row_idx, "detailPemeriksaan_extracted"] = None

    return df


def parse_S_details(df_input):
    df = df_input.copy()

    # make sure the column exists
    if "keluhan_extracted" not in df.columns:
        df["keluhan_extracted"] = None

    row_idx = 0
    row = df.iloc[row_idx]

    raw_text = row.get("detailKeluhan", None)
    if pd.isna(raw_text) or not isinstance(raw_text, str) or raw_text.strip() == "":
        raw_text = row.get("riwayatPenyakit", None)
        if pd.isna(raw_text) or not isinstance(raw_text, str) or raw_text.strip() == "":
            return df

    prompt_text = f"""
        Ekstrak semua keluhan utama secara eksplisit dari teks berikut.
        - Hanya ekstrak yang jelas disebutkan dalam teks.
        - Gunakan format JSON dengan satu key: "keluhan_utama".
        - Nilai dari "keluhan_utama" adalah array berisi string keluhan.
        - Jika tidak ada keluhan, isi array dengan satu item: "Tidak ada keluhan".
        - Jangan tambahkan teks atau penjelasan di luar JSON.

        Teks:
        \"{raw_text.strip()}\" 
    """

    try:
        model = "gemini-2.0-flash"
        contents = [
            types.Content(
                role="user",
                parts=[types.Part.from_text(text=prompt_text)],
            ),
        ]
        generate_content_config = types.GenerateContentConfig(
            response_mime_type="text/plain",
        )

        response_text = ""
        for chunk in client.models.generate_content_stream(
            model=model,
            contents=contents,
            config=generate_content_config,
        ):
            response_text += chunk.text

        clean_text = re.sub(r"```json|```", "", response_text).strip()
        xdata = json.loads(clean_text)

        df.at[row_idx, "keluhan_extracted"] = xdata

    except Exception as e:
        print(f"Row - Error processing column in S: {e}")
        df.at[row_idx, "keluhan_extracted"] = None

    return df


def extract_keluhan_string(x):
    if pd.isna(x):
        return None
    try:
        parsed = ast.literal_eval(x) if isinstance(x, str) else x
        if isinstance(parsed, dict) and 'keluhan_utama' in parsed:
            return ', '.join(parsed['keluhan_utama'])
    except:
        return None


def last_clean_extract(df_input):
    df = df_input.copy()

    # --- Berat Badan ---
    if "beratBadan" in df.columns:
        df["beratBadan"] = (
            df["beratBadan"]
            .astype(str)
            .str.replace(r"[^0-9.,]", "", regex=True)
            .str.replace(",", ".", regex=False)
        )
        df["beratBadan"] = pd.to_numeric(df["beratBadan"], errors="coerce")

    # --- Tinggi Badan ---
    if "tinggiBadan" in df.columns:
        df["tinggiBadan"] = (
            df["tinggiBadan"]
            .astype(str)
            .str.replace(r"[^0-9.,]", "", regex=True)
            .str.replace(",", ".", regex=False)
        )
        df["tinggiBadan"] = pd.to_numeric(df["tinggiBadan"], errors="coerce")

    # --- Tekanan Darah ---
    if "tekananDarah" in df.columns:
        df["tekananDarah"] = (
            df["tekananDarah"]
            .astype(str)
            .str.replace("\\", "/", regex=False)
            .str.replace(r"[^0-9/]", "", regex=True)
        )
        split_bp = df["tekananDarah"].str.split("/", n=1, expand=True)
        df["tekanan_sistolik"] = pd.to_numeric(split_bp[0], errors="coerce")
        df["tekanan_diastolik"] = pd.to_numeric(split_bp[1], errors="coerce")

    # --- Suhu ---
    if "suhu" in df.columns:
        df["suhu"] = (
            df["suhu"]
            .astype(str)
            .str.replace(",", ".", regex=False)
            .str.replace(r"[^0-9.]", "", regex=True)
        )
        df["suhu"] = pd.to_numeric(df["suhu"], errors="coerce")

    # --- Nadi ---
    if "nadi" in df.columns:
        df["nadi"] = (
            df["nadi"]
            .astype(str)
            .str.replace(r"[^0-9]", "", regex=True)
        )
        df["nadi"] = pd.to_numeric(df["nadi"], errors="coerce")

    # --- Pernapasan ---
    if "pernapasan" in df.columns:
        df["pernapasan"] = (
            df["pernapasan"]
            .astype(str)
            .str.replace(r"[^0-9]", "", regex=True)
        )
        df["pernapasan"] = pd.to_numeric(df["pernapasan"], errors="coerce")

    return df


def process_text(text):
    if not text.strip():
        print("Error: Received empty or None text input!")
        return pd.DataFrame()

    df = generate_available_text(text)

    tahun, bulan, hari = zip(*df['pasien.umur'].map(parse_age))
    df['age_days'] = np.array(tahun)*365.25 + np.array(bulan)*30.44 + np.array(hari)
    cut_tuamuda = 49.53 * 365.25            
    bins   = [0, cut_tuamuda, np.inf]    
    labels = ['muda', 'tua']
    df['age_cat'] = pd.cut(df['age_days'], bins=bins, labels=labels, right=True)

    df['pasien.jeniskelamin'] = df['pasien.jeniskelamin'].map({
        'Perempuan': 'F',
        'Laki-Laki': 'M'
    })

    parsed_O_df = parse_O_details(df)
    parsed_S_df = parse_S_details(parsed_O_df)

    parsed_S_df['pengobatan'] = (
        parsed_S_df.get('detailPengobatan', pd.Series([None]*len(parsed_S_df)))
        .combine_first(parsed_S_df.get('hasilInstruksi', pd.Series([None]*len(parsed_S_df))))
    )
    parsed_S_df['keluhan_utama_str'] = parsed_S_df['keluhan_extracted'].apply(extract_keluhan_string)
    
    final_df = last_clean_extract(parsed_S_df)
    return final_df    

def read_docx(file_path):
    doc = Document(file_path)
    return "\n".join([para.text.strip() for para in doc.paragraphs if para.text.strip()])

def read_doc(file_path):
    abs_path = os.path.abspath(file_path)
    doc = Document(abs_path)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text

def select_files():
    return filedialog.askopenfilenames(title="Select Documents", filetypes=[("Word Documents", "*.docx;*.doc")])

def process_selected_files(selected_files):
    df = pd.DataFrame()

    if not selected_files:
        messagebox.showwarning("No Files Selected", "Please select at least one file.")
        return None

    for file_path in selected_files:
        time.sleep(5)
        print("Processing file:", file_path)

        if file_path.endswith(".docx"):
            text = read_docx(file_path)
        elif file_path.endswith(".doc"):
            text = read_doc(file_path)
        else:
            print(f"Unsupported file type: {file_path}")
            continue

        extracted_df = process_text(text)
        df = pd.concat([df, extracted_df], ignore_index=True)

    if df.empty:
        return None  

    return df 

def save_to_excel(df, output_path):
    if not df.empty:
        df.to_excel(output_path, index=True)
        messagebox.showinfo("Success", f"Data saved to {output_path}")
    else:
        messagebox.showwarning("No Data", "No data was extracted to save.")


def extract_llm(selected_files):
    if selected_files:
        df = process_selected_files(selected_files)
        if df is not None:
            output_file = filedialog.asksaveasfilename(
                defaultextension=".xlsx", 
                filetypes=[("Excel Files", "*.xlsx")], 
                title="Save Extracted Data", 
                initialfile="extracted_data.xlsx"
            )
            if output_file:
                save_to_excel(df, output_file)
                return output_file
    return None
